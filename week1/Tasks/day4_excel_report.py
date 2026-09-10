from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Border, Side
import pandas as pd


INPUT_FILE = "../sample_data/invoices.xlsx"
OUTPUT_FILE = "../sample_data/invoice_report.xlsx"


def read_and_filter_invoices():
   #Filter amount above 5000

    df = pd.read_excel(
        INPUT_FILE,
        sheet_name="Sheet1"
    )
   # print("Missing Amount values:")
    df["Amount"].isna()
    df["Amount"] = df["Amount"].fillna(0)

    filtered_df = df[df["Amount"] > 5000]
   

    return df, filtered_df


def calculate_summary(filtered_df):
   #Calculate Total,count,average

    total = filtered_df["Amount"].sum()
    count = len(filtered_df)
    average = filtered_df["Amount"].mean()

    return total, count, average


def highlight_filtered_rows(filtered_df):
   #Highlight the invoices greater than 5000

    workbook = load_workbook(INPUT_FILE)
    worksheet = workbook["Sheet1"]

    yellow_fill = PatternFill(
        fill_type="solid",
        fgColor="FFFFFF00"
    )

    # Find column numbers from the header row
    headers = {
        cell.value: cell.column
        for cell in worksheet[1]
    }

    # Store filtered Invoice IDs for quick lookup
    filtered_ids = set(filtered_df["Invoice ID"])

    # Check each invoice row
    for row in worksheet.iter_rows(min_row=2):

        invoice_id = row[
            headers["Invoice ID"] - 1
        ].value

        if invoice_id in filtered_ids:

            for cell in row:
                cell.fill = yellow_fill

    return workbook


def create_summary_sheet(workbook, total, count, average):
    #Summary Sheet 

    # Remove existing Summary sheet if present
    if "Summary" in workbook.sheetnames:
        del workbook["Summary"]

    summary_ws = workbook.create_sheet("Summary")

    # Merge title cells
    summary_ws.merge_cells("A1:B1")

    summary_ws["A1"] = "Invoice Summary"

    # Format title
    summary_ws["A1"].font = Font(
        bold=True,
        size=14
    )

    # Headers
    summary_ws["A3"] = "Metric"
    summary_ws["B3"] = "Value"

    # Summary values
    summary_ws["A4"] = "Total"
    summary_ws["B4"] = total

    summary_ws["A5"] = "Count"
    summary_ws["B5"] = count

    summary_ws["A6"] = "Average"
    summary_ws["B6"] = average

    # Bold header
    header_font = Font(bold=True)

    summary_ws["A3"].font = header_font
    summary_ws["B3"].font = header_font

    # Create border
    thin_side = Side(style="thin")

    thin_border = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side
    )

    # Apply border to summary table
    for row in summary_ws.iter_rows(
        min_row=3,
        max_row=6,
        min_col=1,
        max_col=2
    ):
        for cell in row:
            cell.border = thin_border

    # Format numbers
    summary_ws["B4"].number_format = '#,##0.00'
    summary_ws["B6"].number_format = '#,##0.00'

    # Set column widths
    summary_ws.column_dimensions["A"].width = 20
    summary_ws.column_dimensions["B"].width = 15

    return summary_ws


def main():


    print("Reading invoice data...")

    # Step 1: Read and filter data
    df, filtered_df = read_and_filter_invoices()

    print("\nFiltered invoices:")
    print(filtered_df)

    # Step 2: Calculate summary
    total, count, average = calculate_summary(
        filtered_df
    )

    print("\nSummary:")
    print("Total:", total)
    print("Count:", count)
    print("Average:", average)

    # Step 3: Highlight filtered rows
    workbook = highlight_filtered_rows(
        filtered_df
    )

    # Step 4: Create Summary sheet
    create_summary_sheet(
        workbook,
        total,
        count,
        average
    )

    # Step 5: Save report
    workbook.save(OUTPUT_FILE)

    print(
        f"\nReport created successfully: {OUTPUT_FILE}"
    )


if __name__ == "__main__":
    main()