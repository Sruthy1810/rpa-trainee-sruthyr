from openpyxl import load_workbook


FILE_PATH = "../sample_data/invoices.xlsx"


def read_invoices(file_path):
    workbook = load_workbook(file_path)
    worksheet = workbook["Sheet1"]

    headers = [
        cell.value
        for cell in worksheet[1]
    ]

    invoices = []

    for row in worksheet.iter_rows(
        min_row=2,
        values_only=True
    ):
        invoice = dict(zip(headers, row))
        invoices.append(invoice)

    return invoices


def main():
    invoices = read_invoices(FILE_PATH)

    for invoice in invoices:
        print(invoice)


if __name__ == "__main__":
    main()